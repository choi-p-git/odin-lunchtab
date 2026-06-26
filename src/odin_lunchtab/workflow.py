from __future__ import annotations

import csv
import io
import logging
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from odin_lunchtab.audit_control import (
    RECONCILIATION_AUDIT_CONTROL_NAME,
    write_reconciliation_audit_control,
)
from odin_lunchtab.profiles import (
    LEGACY_DEFAULT_PROFILE,
    CrosswalkEntry,
    MatchingProfile,
    parse_email_address,
)
from odin_lunchtab.run_reports import (
    RECONCILIATION_RUN_SUMMARY_NAME,
    write_reconciliation_run_summary,
)

ODIN_HEADERS = (
    "Account Type",
    "ID Number",
    "Deposits",
    "Spendings",
    "Balance",
    "Student",
)
PROCESSED_ODIN_HEADERS = (
    "Account Type",
    "ID Number",
    "Deposits",
    "Spendings",
    "Balance",
    "Surname",
    "FirstName",
)
EXCEPTION_HEADERS = (
    "Reason",
    "Account Type",
    "ID Number",
    "Balance",
    "Student",
    "Candidate LoginBarcodes",
    "Candidate Names",
    "Attempted Rules",
    "Transformed Values",
)
AUDIT_HEADERS = (
    "Account Type",
    "Odin ID Number",
    "Odin Student",
    "Lunchtab LoginBarcode",
    "Lunchtab ExternalId",
    "Lunchtab EmailAddress",
    "Lunchtab Name",
    "Profile",
    "Rule",
    "Transformed Value",
    "Matched Target Field",
    "Match Method",
    "OdinBalanceAmount",
)
FINAL_OUTPUT_NAME = "Processed - Odin to Lunchtab Balance Transfer.csv"
EXCEPTIONS_OUTPUT_NAME = "Processed - Odin to Lunchtab Balance Transfer - Exceptions.csv"
MANUAL_REVIEW_EXCEPTIONS_OUTPUT_NAME = (
    "Manual Review Exceptions - Odin to Lunchtab Balance Transfer.csv"
)
MATCH_AUDIT_OUTPUT_NAME = "Accepted Match Audit - Odin to Lunchtab Balance Transfer.csv"
LOGGER = logging.getLogger("odin_lunchtab")


@dataclass(frozen=True)
class CsvEncodingMetadata:
    encoding: str
    used_fallback: bool


@dataclass(frozen=True)
class OdinRecord:
    account_type: str
    id_number: str
    deposits: str
    spendings: str
    balance: str
    student: str
    surname: str
    first_name: str

    def cleaned_row(self) -> dict[str, str]:
        return dict(zip(ODIN_HEADERS, self.raw_values(), strict=True))

    def processed_row(self) -> dict[str, str]:
        values = (
            self.account_type,
            self.id_number,
            self.deposits,
            self.spendings,
            self.balance,
            self.surname,
            self.first_name,
        )
        return dict(zip(PROCESSED_ODIN_HEADERS, values, strict=True))

    def raw_values(self) -> tuple[str, ...]:
        return (
            self.account_type,
            self.id_number,
            self.deposits,
            self.spendings,
            self.balance,
            self.student,
        )


@dataclass(frozen=True)
class MalformedRecord:
    values: tuple[str, ...]
    reason: str


@dataclass(frozen=True)
class OutputPaths:
    cleaned: Path
    processed: Path
    transfer: Path
    exceptions: Path
    manual_review_exceptions: Path
    match_audit: Path | None = None
    audit_control: Path | None = None
    run_summary: Path | None = None
    candidate_matches: Path | None = None


@dataclass(frozen=True)
class RunSummary:
    extracted: int
    malformed: int
    matched_by_id: int
    matched_by_name: int
    exceptions: int
    manual_review_exceptions: int
    lunchtab_rows: int
    output_paths: OutputPaths
    profile_name: str = "Legacy Default"
    profile_schema_version: int = 2
    matches_by_rule: dict[str, int] | None = None
    audit_control_status: str = "PASS"


@dataclass(frozen=True)
class MatchDecision:
    record: OdinRecord
    target_index: int
    method: str
    rule_name: str
    transformed_value: str
    target_field: str = "LoginBarcode"


@dataclass(frozen=True)
class MatchResult:
    output_users: list[dict[str, str]]
    exceptions: list[dict[str, str]]
    counts: Counter[str]
    decisions: list[MatchDecision]


def normalize_name(value: str) -> str:
    decomposed = unicodedata.normalize("NFKD", value.strip())
    without_marks = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(re.findall(r"[a-z0-9]+", without_marks.casefold()))


def split_student_name(student: str) -> tuple[str, str]:
    if "," not in student:
        raise ValueError("student name must use 'Surname, FirstName' format")
    surname, first_name = student.split(",", 1)
    surname = surname.strip()
    first_name = first_name.strip()
    if not surname or not first_name:
        raise ValueError("student surname and first name must both be present")
    return surname, first_name


def first_names_compatible(odin_name: str, lunchtab_name: str) -> bool:
    odin_tokens = normalize_name(odin_name).split()
    lunchtab_tokens = normalize_name(lunchtab_name).split()
    if not odin_tokens or not lunchtab_tokens:
        return False
    odin_first = odin_tokens[0]
    lunchtab_first = lunchtab_tokens[0]
    return (
        odin_first == lunchtab_first
        or odin_first.startswith(lunchtab_first)
        or lunchtab_first.startswith(odin_first)
    )


def names_compatible(record: OdinRecord, user: dict[str, str]) -> bool:
    if normalize_name(record.surname) != normalize_name(user.get("Surname", "")):
        return False
    return first_names_compatible(record.first_name, user.get("FirstName", "")) or (
        bool(user.get("PreferredName", "").strip())
        and first_names_compatible(record.first_name, user["PreferredName"])
    )


def _cell_text(cell: Cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        number_format = str(cell.number_format or "")
        if number_format and set(number_format) == {"0"}:
            return str(value).zfill(len(number_format))
        return str(value)
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return str(value).strip()


def _money_text(value: str) -> str:
    try:
        amount = Decimal(value.replace(",", "").replace("$", "").strip())
    except (InvalidOperation, AttributeError) as error:
        raise ValueError("invalid monetary value") from error
    if not amount.is_finite():
        raise ValueError("invalid monetary value")
    return format(amount, "f")


def extract_odin_report(path: Path) -> tuple[list[OdinRecord], list[MalformedRecord]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            valid: list[OdinRecord] = []
            malformed: list[MalformedRecord] = []
            found_headers = False
            for row in sheet.iter_rows():
                values = tuple(_cell_text(cell) for cell in row[: len(ODIN_HEADERS)])
                if not found_headers:
                    found_headers = values == ODIN_HEADERS
                    continue
                if not any(values):
                    continue
                if values[1] == "Patron Count:" or values[0] == "Account Code":
                    break
                if not values[0] and values[1].startswith(("Total ", "Balances By ")):
                    break
                if not (values[0] or values[1]):
                    continue
                try:
                    if not all(values):
                        raise ValueError("missing one or more required Odin fields")
                    surname, first_name = split_student_name(values[5])
                    deposits, spendings, balance = (
                        _money_text(values[2]),
                        _money_text(values[3]),
                        _money_text(values[4]),
                    )
                except ValueError as error:
                    malformed.append(MalformedRecord(values=values, reason=str(error)))
                    continue
                valid.append(
                    OdinRecord(
                        account_type=values[0],
                        id_number=values[1],
                        deposits=deposits,
                        spendings=spendings,
                        balance=balance,
                        student=values[5],
                        surname=surname,
                        first_name=first_name,
                    )
                )
            if found_headers:
                return valid, malformed
    finally:
        workbook.close()
    raise ValueError(f"Could not find the required Odin headers in {path}")


def _decode_csv_bytes(path: Path) -> tuple[str, CsvEncodingMetadata]:
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        raise ValueError("CSV uses unsupported UTF-32 encoding.")
    if data.startswith(b"\xef\xbb\xbf"):
        candidates = (("utf-8-sig", "UTF-8 with BOM", False),)
    elif data.startswith(b"\xff\xfe"):
        candidates = (("utf-16", "UTF-16 LE with BOM", False),)
    elif data.startswith(b"\xfe\xff"):
        candidates = (("utf-16", "UTF-16 BE with BOM", False),)
    else:
        if b"\x00" in data:
            raise ValueError("CSV contains NUL bytes or unsupported UTF-16 without a BOM.")
        candidates = (
            ("utf-8", "UTF-8", False),
            ("cp1252", "Windows-1252", True),
        )

    last_error: UnicodeDecodeError | None = None
    for codec, label, used_fallback in candidates:
        try:
            text = data.decode(codec, errors="strict")
            metadata = CsvEncodingMetadata(label, used_fallback)
            LOGGER.info(
                "CSV decoded: encoding=%s fallback=%s",
                label,
                used_fallback,
            )
            return text, metadata
        except UnicodeDecodeError as error:
            last_error = error
    raise ValueError(
        "CSV encoding is unsupported or contains invalid text. "
        "Supported encodings are UTF-8, Windows-1252, and BOM-marked UTF-16."
    ) from last_error


def _validate_csv_text(text: str) -> None:
    if "\ufffd" in text:
        raise ValueError("CSV contains Unicode replacement characters and may be corrupted.")
    unsafe = {
        character
        for character in text
        if (ord(character) < 32 and character not in "\t\r\n") or 0x7F <= ord(character) <= 0x9F
    }
    if unsafe:
        raise ValueError("CSV contains unsupported control characters or binary data.")


def read_csv_with_metadata(
    path: Path,
) -> tuple[list[str], list[dict[str, str]], CsvEncodingMetadata]:
    text, metadata = _decode_csv_bytes(path)
    _validate_csv_text(text)
    try:
        reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
        if reader.fieldnames is None:
            raise ValueError(f"CSV has no header row: {path}")
        return list(reader.fieldnames), list(reader), metadata
    except csv.Error as error:
        raise ValueError(f"CSV structure is malformed: {error}") from error


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    headers, rows, _ = read_csv_with_metadata(path)
    return headers, rows


def _discover_odin(raw_data_dir: Path) -> Path:
    candidates = [
        path
        for path in raw_data_dir.glob("*.xlsx")
        if not path.name.startswith("~$") and path.is_file()
    ]
    if len(candidates) != 1:
        raise ValueError(
            f"Expected exactly one non-temporary Odin .xlsx in {raw_data_dir}; "
            f"found {len(candidates)}"
        )
    return candidates[0]


def _discover_lunchtab(raw_data_dir: Path) -> Path:
    exact = raw_data_dir / "Lunchtab Users.csv"
    if exact.is_file():
        return exact
    candidates = [path for path in raw_data_dir.glob("*.csv") if path.is_file()]
    if len(candidates) != 1:
        raise ValueError(
            f"Expected Lunchtab Users.csv or exactly one .csv in {raw_data_dir}; "
            f"found {len(candidates)}"
        )
    return candidates[0]


def _output_paths(output_dir: Path, odin_path: Path) -> OutputPaths:
    return OutputPaths(
        cleaned=output_dir / f"cleaned - {odin_path.stem}.csv",
        processed=output_dir / f"processed - {odin_path.stem}.csv",
        transfer=output_dir / FINAL_OUTPUT_NAME,
        exceptions=output_dir / EXCEPTIONS_OUTPUT_NAME,
        manual_review_exceptions=output_dir / MANUAL_REVIEW_EXCEPTIONS_OUTPUT_NAME,
        match_audit=output_dir / MATCH_AUDIT_OUTPUT_NAME,
        audit_control=output_dir / RECONCILIATION_AUDIT_CONTROL_NAME,
        run_summary=output_dir / RECONCILIATION_RUN_SUMMARY_NAME,
        candidate_matches=output_dir / "Manual Review Candidate Matches.csv",
    )


def _write_csv(path: Path, fieldnames: Iterable[str], rows: Iterable[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _candidate_text(candidates: Iterable[int], users: list[dict[str, str]]) -> tuple[str, str]:
    candidate_rows = [users[index] for index in candidates]
    barcodes = " | ".join(row.get("LoginBarcode", "") for row in candidate_rows)
    names = " | ".join(
        f"{row.get('Surname', '')}, {row.get('FirstName', '')}".strip() for row in candidate_rows
    )
    return barcodes, names


def _exception_row(
    reason: str,
    record: OdinRecord | MalformedRecord,
    users: list[dict[str, str]],
    candidates: Iterable[int] = (),
    attempted_rules: Iterable[str] = (),
    transformed_values: Iterable[str] = (),
) -> dict[str, str]:
    if isinstance(record, OdinRecord):
        account_type = record.account_type
        id_number = record.id_number
        balance = record.balance
        student = record.student
    else:
        padded = (*record.values, *([""] * (len(ODIN_HEADERS) - len(record.values))))
        account_type, id_number, _, _, balance, student = padded[: len(ODIN_HEADERS)]
    barcodes, names = _candidate_text(candidates, users)
    return {
        "Reason": reason,
        "Account Type": account_type,
        "ID Number": id_number,
        "Balance": balance,
        "Student": student,
        "Candidate LoginBarcodes": barcodes,
        "Candidate Names": names,
        "Attempted Rules": " | ".join(attempted_rules),
        "Transformed Values": " | ".join(transformed_values),
    }


def _name_candidates(record: OdinRecord, users: list[dict[str, str]]) -> list[int]:
    return [index for index, user in enumerate(users) if names_compatible(record, user)]


def _target_indexes(
    users: list[dict[str, str]],
) -> dict[str, defaultdict[str, list[int]]]:
    indexes = {
        "LoginBarcode": defaultdict(list),
        "ExternalId": defaultdict(list),
        "EmailUsername": defaultdict(list),
    }
    for index, user in enumerate(users):
        for field in ("LoginBarcode", "ExternalId"):
            indexes[field][user.get(field, "").strip()].append(index)
        parsed = parse_email_address(user.get("EmailAddress", ""))
        if parsed:
            username, _ = parsed
            indexes["EmailUsername"][username.casefold()].append(index)
    return indexes


def _lookup_value(target_field: str, value: str) -> str:
    return value.casefold() if target_field == "EmailUsername" else value


def _email_candidate_allowed(
    user: dict[str, str],
    allowed_domains: tuple[str, ...],
) -> bool:
    if not allowed_domains:
        return True
    parsed = parse_email_address(user.get("EmailAddress", ""))
    allowed = {domain.casefold() for domain in allowed_domains}
    return bool(parsed and parsed[1] in allowed)


def _crosswalk_by_source(
    entries: tuple[CrosswalkEntry, ...],
) -> dict[tuple[str, str], CrosswalkEntry]:
    return {(entry.account_type, entry.odin_id): entry for entry in entries}


def _audit_row(
    decision: MatchDecision,
    user: dict[str, str],
    profile: MatchingProfile,
) -> dict[str, str]:
    return {
        "Account Type": decision.record.account_type,
        "Odin ID Number": decision.record.id_number,
        "Odin Student": decision.record.student,
        "Lunchtab LoginBarcode": user.get("LoginBarcode", ""),
        "Lunchtab ExternalId": user.get("ExternalId", ""),
        "Lunchtab EmailAddress": user.get("EmailAddress", ""),
        "Lunchtab Name": f"{user.get('Surname', '')}, {user.get('FirstName', '')}".strip(),
        "Profile": profile.name,
        "Rule": decision.rule_name,
        "Transformed Value": decision.transformed_value,
        "Matched Target Field": decision.target_field,
        "Match Method": decision.method,
        "OdinBalanceAmount": decision.record.balance,
    }


def match_balances_detailed(
    records: list[OdinRecord],
    users: list[dict[str, str]],
    malformed: list[MalformedRecord] | None = None,
    profile: MatchingProfile | None = None,
) -> MatchResult:
    profile = profile or LEGACY_DEFAULT_PROFILE
    profile.validate()
    indexes = _target_indexes(users)
    crosswalk = _crosswalk_by_source(profile.crosswalk)
    odin_id_counts = Counter((record.account_type, record.id_number) for record in records)
    proposed: list[MatchDecision] = []
    exceptions = [_exception_row(item.reason, item, users) for item in (malformed or [])]

    for record in records:
        source_key = (record.account_type, record.id_number)
        if odin_id_counts[source_key] > 1:
            exceptions.append(_exception_row("duplicate Odin ID Number", record, users))
            continue

        mapping = crosswalk.get(source_key)
        if mapping is not None:
            candidates = [
                candidate
                for candidate in indexes[mapping.target_field].get(
                    _lookup_value(mapping.target_field, mapping.target_value),
                    [],
                )
                if _email_candidate_allowed(users[candidate], mapping.allowed_email_domains)
            ]
            if len(candidates) != 1:
                reason = (
                    "stale crosswalk target"
                    if not candidates
                    else f"duplicate Lunchtab {mapping.target_field}"
                )
                exceptions.append(
                    _exception_row(
                        reason,
                        record,
                        users,
                        candidates,
                        ["Crosswalk"],
                        [mapping.target_value],
                    )
                )
            elif names_compatible(record, users[candidates[0]]):
                proposed.append(
                    MatchDecision(
                        record,
                        candidates[0],
                        "crosswalk",
                        "Crosswalk",
                        mapping.target_value,
                        mapping.target_field,
                    )
                )
            else:
                exceptions.append(
                    _exception_row(
                        "crosswalk matched but name validation failed",
                        record,
                        users,
                        candidates,
                        ["Crosswalk"],
                        [mapping.target_value],
                    )
                )
            continue

        applicable = [rule for rule in profile.rules if rule.applies_to(record.account_type)]
        legacy_exact = (
            profile.profile_id == LEGACY_DEFAULT_PROFILE.profile_id
            and len(applicable) == 1
            and applicable[0].name == "Exact LoginBarcode"
        )
        attempted_rules: list[str] = []
        transformed_values: list[str] = []
        valid_targets: defaultdict[int, list[tuple[str, str, str]]] = defaultdict(list)
        duplicate_candidates: list[int] = []
        name_failed_candidates: list[int] = []
        for rule in applicable:
            transformed = rule.transform(record.id_number)
            attempted_rules.append(rule.name)
            transformed_values.append(transformed)
            candidates = [
                candidate
                for candidate in indexes[rule.target_field].get(
                    _lookup_value(rule.target_field, transformed),
                    [],
                )
                if _email_candidate_allowed(users[candidate], rule.allowed_email_domains)
            ]
            if len(candidates) > 1:
                duplicate_candidates.extend(candidates)
            elif len(candidates) == 1:
                candidate = candidates[0]
                if names_compatible(record, users[candidate]):
                    valid_targets[candidate].append((rule.name, transformed, rule.target_field))
                else:
                    name_failed_candidates.append(candidate)

        if duplicate_candidates:
            exceptions.append(
                _exception_row(
                    "duplicate Lunchtab LoginBarcode"
                    if legacy_exact
                    else "duplicate Lunchtab identifier",
                    record,
                    users,
                    duplicate_candidates,
                    attempted_rules,
                    transformed_values,
                )
            )
            continue
        if len(valid_targets) > 1:
            exceptions.append(
                _exception_row(
                    "matching rules resolved to different Lunchtab users",
                    record,
                    users,
                    valid_targets.keys(),
                    attempted_rules,
                    transformed_values,
                )
            )
            continue
        if len(valid_targets) == 1:
            target = next(iter(valid_targets))
            matches = valid_targets[target]
            rule_names = " + ".join(rule_name for rule_name, _, _ in matches)
            transformed = " + ".join(value for _, value, _ in matches)
            target_fields = " + ".join(field for _, _, field in matches)
            legacy_id = (
                profile.profile_id == LEGACY_DEFAULT_PROFILE.profile_id
                and rule_names == "Exact LoginBarcode"
            )
            proposed.append(
                MatchDecision(
                    record,
                    target,
                    "id" if legacy_id else "rule",
                    rule_names,
                    transformed,
                    target_fields,
                )
            )
            continue
        if name_failed_candidates:
            exceptions.append(
                _exception_row(
                    "barcode matched but name validation failed"
                    if legacy_exact
                    else "identifier matched but name validation failed",
                    record,
                    users,
                    name_failed_candidates,
                    attempted_rules,
                    transformed_values,
                )
            )
            continue

        if profile.allows_name_fallback(record.account_type):
            name_candidates = _name_candidates(record, users)
            if len(name_candidates) == 1:
                proposed.append(
                    MatchDecision(
                        record,
                        name_candidates[0],
                        "name",
                        "Unique name fallback",
                        "",
                        "Name",
                    )
                )
            elif len(name_candidates) > 1:
                exceptions.append(
                    _exception_row(
                        "ambiguous name match",
                        record,
                        users,
                        name_candidates,
                        attempted_rules,
                        transformed_values,
                    )
                )
            else:
                exceptions.append(
                    _exception_row(
                        "no Lunchtab match",
                        record,
                        users,
                        attempted_rules=attempted_rules,
                        transformed_values=transformed_values,
                    )
                )
        else:
            exceptions.append(
                _exception_row(
                    "no configured identifier match",
                    record,
                    users,
                    attempted_rules=attempted_rules,
                    transformed_values=transformed_values,
                )
            )

    target_counts = Counter(decision.target_index for decision in proposed)
    accepted: list[MatchDecision] = []
    for decision in proposed:
        if target_counts[decision.target_index] > 1:
            exceptions.append(
                _exception_row(
                    "multiple Odin records matched one Lunchtab user",
                    decision.record,
                    users,
                    [decision.target_index],
                    [decision.rule_name],
                    [decision.transformed_value],
                )
            )
        else:
            accepted.append(decision)

    output_users = [dict(user) for user in users]
    for user in output_users:
        user["OdinBalanceAmount"] = ""
    counts: Counter[str] = Counter()
    for decision in accepted:
        output_users[decision.target_index]["OdinBalanceAmount"] = decision.record.balance
        counts[decision.method] += 1
        if decision.method == "rule":
            counts[f"rule:{decision.rule_name}"] += 1
        elif decision.method == "crosswalk":
            counts["rule:Crosswalk"] += 1
    return MatchResult(output_users, exceptions, counts, accepted)


def match_balances(
    records: list[OdinRecord],
    users: list[dict[str, str]],
    malformed: list[MalformedRecord] | None = None,
    profile: MatchingProfile | None = None,
) -> tuple[list[dict[str, str]], list[dict[str, str]], Counter[str]]:
    result = match_balances_detailed(records, users, malformed, profile)
    return result.output_users, result.exceptions, result.counts


def run_workflow(
    *,
    raw_data_dir: Path,
    output_dir: Path,
    odin_path: Path | None = None,
    lunchtab_path: Path | None = None,
    overwrite: bool = False,
    profile: MatchingProfile | None = None,
) -> RunSummary:
    profile = profile or LEGACY_DEFAULT_PROFILE
    profile.validate()
    odin_path = odin_path or _discover_odin(raw_data_dir)
    lunchtab_path = lunchtab_path or _discover_lunchtab(raw_data_dir)
    if not odin_path.is_file():
        raise FileNotFoundError(f"Odin report does not exist: {odin_path}")
    if not lunchtab_path.is_file():
        raise FileNotFoundError(f"Lunchtab export does not exist: {lunchtab_path}")

    paths = _output_paths(output_dir, odin_path)
    existing = [path for path in paths.__dict__.values() if path.exists()]
    if existing and not overwrite:
        names = ", ".join(str(path) for path in existing)
        raise FileExistsError(f"Refusing to overwrite existing output(s): {names}")

    records, malformed = extract_odin_report(odin_path)
    lunchtab_headers, users = read_csv(lunchtab_path)
    required_headers = {
        "FirstName",
        "PreferredName",
        "Surname",
        "LoginBarcode",
        "DefaultFamilyCode",
    }
    missing = sorted(required_headers - set(lunchtab_headers))
    profile_fields = {
        "EmailAddress" if rule.target_field == "EmailUsername" else rule.target_field
        for rule in profile.rules
        if rule.enabled
    } | {
        "EmailAddress" if entry.target_field == "EmailUsername" else entry.target_field
        for entry in profile.crosswalk
    }
    missing = sorted(set(missing) | (profile_fields - set(lunchtab_headers)))
    if missing:
        raise ValueError(f"Lunchtab export is missing required columns: {', '.join(missing)}")
    if "DefaultFamilyBalanceAmount" not in lunchtab_headers:
        raise ValueError("Lunchtab export is missing DefaultFamilyBalanceAmount")

    result = match_balances_detailed(records, users, malformed, profile)
    output_users, exceptions, counts = (
        result.output_users,
        result.exceptions,
        result.counts,
    )
    final_headers = list(lunchtab_headers)
    insertion_index = final_headers.index("DefaultFamilyBalanceAmount") + 1
    final_headers.insert(insertion_index, "OdinBalanceAmount")

    output_dir.mkdir(parents=True, exist_ok=True)
    _write_csv(paths.cleaned, ODIN_HEADERS, (record.cleaned_row() for record in records))
    _write_csv(
        paths.processed,
        PROCESSED_ODIN_HEADERS,
        (record.processed_row() for record in records),
    )
    _write_csv(paths.transfer, final_headers, output_users)
    _write_csv(paths.exceptions, EXCEPTION_HEADERS, exceptions)
    manual_review_exceptions = [
        exception for exception in exceptions if exception["Reason"] != "no Lunchtab match"
    ]
    _write_csv(
        paths.manual_review_exceptions,
        EXCEPTION_HEADERS,
        manual_review_exceptions,
    )
    from odin_lunchtab.exception_candidates import write_manual_review_candidate_report

    write_manual_review_candidate_report(
        manual_review_exceptions_path=paths.manual_review_exceptions,
        lunchtab_path=lunchtab_path,
        output_dir=output_dir,
    )
    if paths.match_audit is not None:
        _write_csv(
            paths.match_audit,
            AUDIT_HEADERS,
            (
                _audit_row(decision, users[decision.target_index], profile)
                for decision in result.decisions
            ),
        )
    audit_control_status = "PASS"
    if paths.audit_control is not None and paths.match_audit is not None:
        audit_control_status = write_reconciliation_audit_control(
            path=paths.audit_control,
            records=records,
            malformed=malformed,
            decisions=result.decisions,
            exceptions=exceptions,
            processed_odin_name=paths.processed.name,
            match_audit_name=paths.match_audit.name,
            exceptions_name=paths.exceptions.name,
        )
        if audit_control_status != "PASS":
            raise RuntimeError("Reconciliation audit-control totals failed the integrity audit.")
    if paths.run_summary is not None and paths.audit_control is not None:
        write_reconciliation_run_summary(
            path=paths.run_summary,
            profile_name=profile.name,
            valid_odin_rows=len(records),
            malformed_rows=len(malformed),
            matched_by_id=counts["id"] + counts["rule"] + counts["crosswalk"],
            matched_by_name=counts["name"],
            exceptions=len(exceptions),
            manual_review_exceptions=len(manual_review_exceptions),
            lunchtab_rows=len(users),
            audit_control_status=audit_control_status,
            audit_control_name=paths.audit_control.name,
            manual_review_name=paths.manual_review_exceptions.name,
            exceptions_name=paths.exceptions.name,
            transfer_name=paths.transfer.name,
        )

    return RunSummary(
        extracted=len(records),
        malformed=len(malformed),
        matched_by_id=counts["id"] + counts["rule"] + counts["crosswalk"],
        matched_by_name=counts["name"],
        exceptions=len(exceptions),
        manual_review_exceptions=len(manual_review_exceptions),
        lunchtab_rows=len(users),
        output_paths=paths,
        profile_name=profile.name,
        profile_schema_version=profile.schema_version,
        matches_by_rule={
            key.removeprefix("rule:"): value
            for key, value in counts.items()
            if key.startswith("rule:")
        },
        audit_control_status=audit_control_status,
    )
